#!/usr/bin/env python3
"""スライド画像の一覧から 16:9 の .pptx を組み立てる。

    python3 _pptx_from_images.py <画像一覧ファイル> <出力先.pptx>

build_pptx.mjs から呼ばれる補助スクリプト。単体で使うことは想定していない。
1 枚の画像をスライド全面に敷くだけなので、HTML と見た目が完全に一致する。
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

# 16:9。deck.html の @page（13.333in x 7.5in）と揃える。
# EMU は 1 inch = 914400。
SLIDE_W = Emu(int(13.333 * 914400))
SLIDE_H = Emu(int(7.5 * 914400))


def main() -> int:
    if len(sys.argv) != 3:
        print("使い方: _pptx_from_images.py <画像一覧ファイル> <出力先.pptx>", file=sys.stderr)
        return 2

    list_file, out = Path(sys.argv[1]), Path(sys.argv[2])
    images = [Path(p) for p in list_file.read_text(encoding="utf-8").splitlines() if p.strip()]
    if not images:
        print("画像が 1 枚もありません", file=sys.stderr)
        return 1

    missing = [p for p in images if not p.exists()]
    if missing:
        print(f"画像が見つかりません: {missing[0]}", file=sys.stderr)
        return 1

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # レイアウト6 = 白紙。既定のタイトル枠が乗ると画像の上に空の枠が重なる。
    blank = prs.slide_layouts[6]
    for img in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0, width=SLIDE_W, height=SLIDE_H)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"  {len(images)} 枚を pptx に格納しました")
    return 0


if __name__ == "__main__":
    sys.exit(main())
